import io

from conftest import case_id, client, hdr


def analyze(user, files):
    return client.post("/api/v1/intake/analyze", headers=hdr(user), files=[("files", f) for f in files])


def docx_bytes(paragraphs, table=None):
    from docx import Document
    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for i, row in enumerate(table):
            for j, cell in enumerate(row):
                t.cell(i, j).text = cell
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


NEW_FIR = ("FIRST INFORMATION REPORT (EXCISE)\nCr.No.: 301/2026        Date of registration: 05.10.2026\n"
           "On 05.10.2026 at Anakapalle, Ravi Kumar, S/o Subba Rao, aged 30, R/o Paderu, was found with 12 bottles of NDPL in AP 31 TX 9999. "
           "Offence under: AP Excise Act Sec 34(a).")


def test_analyze_reads_type_case_ref_and_date_but_stores_nothing():
    before = len(client.get("/api/v1/cases", headers=hdr("inv_vizag")).json())
    r = analyze("inv_vizag", [("fir.txt", NEW_FIR.encode(), "text/plain")]).json()[0]
    assert r["ok"] and r["doc_type"] == "fir" and r["case_ref"] == "301/2026" and r["registered_on"] == "2026-10-05"
    assert r["matches"] == [] and r["matched_case_id"] is None  # new case: nothing to match
    assert len(client.get("/api/v1/cases", headers=hdr("inv_vizag")).json()) == before


def test_matches_an_existing_case_only_within_the_users_scope():
    text = "SEIZURE PANCHANAMA\nCr.No.: 041/2025 seized 5 bottles".encode()
    own = analyze("inv_vizag", [("p.txt", text)]).json()[0]
    assert own["doc_type"] == "seizure" and own["matched_case_id"] == case_id("EXC/NRP/041/2025")
    guntur_ref = "FIRST INFORMATION REPORT\nCr.No.: 45/2025".encode()
    assert analyze("inv_vizag", [("g.txt", guntur_ref)]).json()[0]["matches"] == []    # another district: not suggested
    assert analyze("inv_guntur", [("g.txt", guntur_ref)]).json()[0]["matched_case_id"] == case_id("Cr.45/2025-TNL")


def test_word_documents_are_read_including_tables():
    data = docx_bytes(["FIRST INFORMATION REPORT (EXCISE)", "Cr.No.: 302/2026"], [["Accused", "Mohan Rao, S/o Narayana, aged 41, R/o Paderu"]])
    r = analyze("inv_vizag", [("fir.docx", data)]).json()[0]
    assert r["ok"] and r["doc_type"] == "fir" and r["case_ref"] == "302/2026"


def test_bad_files_are_reported_per_file_and_do_not_block_others():
    res = analyze("inv_vizag", [("ok.txt", NEW_FIR.encode()), ("x.exe", b"MZ"), ("bad.docx", b"not a docx")]).json()
    assert [r["ok"] for r in res] == [True, False, False]
    assert "Accepted" in res[1]["error"] and "could not be read" in res[2]["error"]


def test_duplicates_are_flagged_and_roles_enforced():
    text = (NEW_FIR + " duplicate probe").encode()
    cid = case_id("EXC/NRP/041/2025")
    client.post("/api/v1/ingest", headers=hdr("inv_vizag"), data={"case_id": cid}, files={"file": ("a.txt", text)})
    assert analyze("inv_vizag", [("a.txt", text)]).json()[0]["duplicate_of"]["case_no"] == "EXC/NRP/041/2025"
    assert analyze("inv_guntur", [("a.txt", text)]).json()[0]["duplicate_of"] is None  # cannot see other district's documents
    assert analyze("analyst_vizag", [("a.txt", text)]).status_code == 403


def test_create_case_validates_scope_and_uniqueness_then_ingests_end_to_end():
    body = {"case_no": "EXC/ANK/301/2026", "district": "Visakhapatnam", "station": "Excise Station, Anakapalle", "registered_on": "2026-10-05"}
    assert client.post("/api/v1/cases", headers=hdr("analyst_vizag"), json=body).status_code == 403
    assert client.post("/api/v1/cases", headers=hdr("inv_guntur"), json=body).status_code == 403   # not their district
    assert client.post("/api/v1/cases", headers=hdr("inv_vizag"), json=body | {"registered_on": "05-10-2026"}).status_code == 422
    r = client.post("/api/v1/cases", headers=hdr("inv_vizag"), json=body)
    assert r.status_code == 201
    assert client.post("/api/v1/cases", headers=hdr("inv_vizag"), json=body).status_code == 409
    cid = r.json()["id"]
    up = client.post("/api/v1/ingest", headers=hdr("inv_vizag"), data={"case_id": cid, "doc_type": "fir"}, files={"file": ("fir.docx", docx_bytes(NEW_FIR.split("\n")))})
    assert up.status_code == 202
    facts = client.get(f"/api/v1/cases/{cid}/facts", headers=hdr("inv_vizag")).json()
    assert {"Ravi Kumar", "AP 31 TX 9999"} <= {f["original"] for f in facts}
    assert client.get(f"/api/v1/cases/{cid}", headers=hdr("inv_guntur")).status_code == 404  # new case is scoped too
    assert any(a["action"] == "create_case" for a in client.get("/api/v1/audit", headers=hdr("sup_vizag")).json())
